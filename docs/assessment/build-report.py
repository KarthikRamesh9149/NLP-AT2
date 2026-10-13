from pathlib import Path
from datetime import datetime, timezone
import re, json, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from pdf2image import convert_from_path

ROOT = Path(__file__).resolve().parent
REPO = next((p for p in ROOT.parents if (p/'evaluation/results').is_dir()), ROOT.parent.parent/'formbridge')
RESULTS = REPO / 'evaluation/results'
DEV = json.loads((RESULTS/'sgd-dev-20-report-v18-20261013-development.json').read_text())
REG = json.loads((RESULTS/'sgd-test-20-report-v18-20261013-regression.json').read_text())
values = {
 'DEV_RULES': round(DEV['summary']['rules']['canonicalFieldAccuracy']*100),
 'DEV_DIRECT': round(DEV['summary']['model']['canonicalFieldAccuracy']*100),
 'DEV_FB': round(DEV['summary']['formbridge']['canonicalFieldAccuracy']*100),
 'DEV_COMPLETE': round(DEV['summary']['formbridge']['canonicalCompleteFormAccuracy']*20),
 'REG_FB': round(REG['summary']['formbridge']['canonicalFieldAccuracy']*100),
 'COST_ABSTRACT': 'The higher development score required more provider calls and approximately twice the mean dialogue replay time of direct extraction.'
}
source = (ROOT/'report-template.md').read_text()
for key, value in values.items(): source = source.replace('{{'+key+'}}',str(value))
assert '{{' not in source
(ROOT/'FormBridge-AT2-Report.md').write_text(source)

# Explanatory architecture figure using ordinary vector drawing tools.
asset_pdf=ROOT/'assets/architecture-current.pdf'
c=canvas.Canvas(str(asset_pdf),pagesize=(720,238))
c.setFillColor(HexColor('#ffffff'));c.rect(0,0,720,238,fill=1,stroke=0)
def box(x,y,w,h,title,detail):
 c.setFillColor(HexColor('#f5f5f3'));c.setStrokeColor(HexColor('#333333'));c.setLineWidth(.7);c.rect(x,y,w,h,fill=1,stroke=1)
 c.setFillColor(HexColor('#111111'));c.setFont('Helvetica-Bold',10);c.drawCentredString(x+w/2,y+h*.68,title)
 c.setFont('Helvetica',9);c.setFillColor(HexColor('#333333'));c.drawCentredString(x+w/2,y+h*.29,detail)
def arrow(x1,y1,x2,y2):
 c.setStrokeColor(HexColor('#444444'));c.setFillColor(HexColor('#444444'));c.setLineWidth(.8);c.line(x1,y1,x2,y2)
 if x1==x2:points=[(x2,y2),(x2-3,y2+6),(x2+3,y2+6)]
 elif x1<x2:points=[(x2,y2),(x2-6,y2-3),(x2-6,y2+3)]
 else:points=[(x2,y2),(x2+6,y2-3),(x2+6,y2+3)]
 p=c.beginPath();p.moveTo(*points[0]);p.lineTo(*points[1]);p.lineTo(*points[2]);p.close();c.drawPath(p,fill=1,stroke=0)
box(4,148,161,62,'Conversation input','Text / reviewed voice / guest page')
box(184,148,158,62,'Server boundary','Schema / identity / allowance')
box(361,148,153,62,'OpenAI extraction','Values and quotations')
box(533,148,183,62,'Validation and state','Evidence / constraints / revisions')
arrow(165,179,184,179);arrow(342,179,361,179);arrow(514,179,533,179)
box(533,22,183,67,'Staff review','Sources / edits / approval')
box(312,22,183,67,'Private team records','Supabase / roles / version checks')
box(14,22,245,67,'Guest and handoff views','Limited review links / simulated providers')
arrow(624,148,624,89);arrow(533,55,495,55);arrow(312,55,259,55)
c.setFillColor(HexColor('#555555'));c.setFont('Helvetica',9);c.drawCentredString(360,226,'INPUT  /  LANGUAGE PROCESSING  /  REVIEWED ENQUIRY')
c.showPage();c.save()
convert_from_path(str(asset_pdf),dpi=220)[0].save(ROOT/'assets/architecture-current.png')

# The complete rewrite retains the reference page system and cover direction.
doc = Document(ROOT/'report-base.docx')
body = doc._element.body
for element in list(body):
 if element.tag != qn('w:sectPr'): body.remove(element)
sec=doc.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.68)
sec.left_margin=Inches(.76);sec.right_margin=Inches(.76)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
sec.different_first_page_header_footer=True
for container in [sec.header,sec.footer,sec.first_page_header,sec.first_page_footer]:
 for e in list(container._element):container._element.remove(e)
styles=doc.styles
n=styles['Normal'];n.font.name='Arial';n.font.size=Pt(11);n.font.color.rgb=RGBColor.from_string('151515')
n.paragraph_format.line_spacing=1.1;n.paragraph_format.space_before=Pt(0);n.paragraph_format.space_after=Pt(6)
n.paragraph_format.widow_control=True
for name,size in [('Heading 1',18),('Heading 2',12)]:
 s=styles[name];s.font.name='Georgia' if name=='Heading 1' else 'Arial';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor.from_string('111111')
 s.paragraph_format.space_before=Pt(9 if name=='Heading 2' else 0);s.paragraph_format.space_after=Pt(7)
 s.paragraph_format.line_spacing=1.04;s.paragraph_format.keep_with_next=True
s=styles['Caption'];s.font.name='Arial';s.font.size=Pt(9.5);s.font.italic=False;s.font.color.rgb=RGBColor.from_string('454545')
s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.05
h=sec.header.add_paragraph('FormBridge  |  36118 Applied Natural Language Processing')
h.paragraph_format.space_after=Pt(0)
for r in h.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
ft=sec.footer.add_paragraph();ft.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');ft._p.append(field)
for r in ft.runs:r.font.size=Pt(9)
doc.core_properties.title='FormBridge Restaurant Enquiry Workspace'
doc.core_properties.subject='36118 Applied Natural Language Processing: Assessment Task 2'
doc.core_properties.author='FormBridge project group'
doc.core_properties.keywords='dialogue state tracking, restaurant enquiries, natural language processing'
doc.core_properties.comments=''
doc.core_properties.created=datetime.now(timezone.utc)
doc.core_properties.modified=doc.core_properties.created

WIDTH=6.7477

def link(p,label,url):
 rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'173D68');pr.append(color);r.append(pr)
 t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def rich(p,text):
 for part in re.split(r'(\*\*.*?\*\*|\*[^*]+\*|`[^`]+`|https?://[^\s]+)',text):
  if not part:continue
  if part.startswith('http'):link(p,part,part);continue
  strong=part.startswith('**') and part.endswith('**');italic=not strong and part.startswith('*') and part.endswith('*');code=part.startswith('`') and part.endswith('`')
  value=part[2:-2] if strong else part[1:-1] if italic or code else part
  r=p.add_run(value);r.bold=strong;r.italic=italic
  if code:r.font.name='Liberation Mono';r.font.size=Pt(9.5)

def para(text,page):
 caption=text.startswith('*Figure ') or text.startswith('*Table ')
 p=doc.add_paragraph(style='Caption' if caption else None)
 if caption:
  is_table=text.startswith('*Table ')
  text=text[1:-1] if text.endswith('*') else text[1:]
  if is_table:p.paragraph_format.keep_with_next=True
 rich(p,text)
 if page==15:
  p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)
  p.paragraph_format.line_spacing=2;p.paragraph_format.space_after=Pt(0)
 return p

def table(lines,page):
 rows=[[x.strip() for x in line.strip().strip('|').split('|')] for line in lines]
 rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',x) for x in r)]
 cols=len(rows[0]);t=doc.add_table(rows=0,cols=cols);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths={2:[2.2,4.5477],3:[1.6,2.45,2.6977],4:[1.95,1.8,1.5,1.4977],5:[1.55,1.3,1.3,1.3,1.2977],6:[2.0,.95,.95,.95,.95,.9477]}[cols]
 if page==2:widths=[5.9477,.8]
 if page==10:widths=[1.28,1.67,3.7977]
 if page==12:widths=[1.3,2.88,2.5677]
 if page==18:widths=[3.55,3.1977]
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 for ri,row in enumerate(rows):
  cells=t.add_row().cells
  for ci,txt in enumerate(row):
   c=cells[ci];c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   p=c.paragraphs[0];p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.05
   rich(p,txt)
   for r in p.runs:
    r.font.size=Pt(9.5 if cols>4 else 10)
    if ri==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
   pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'333333' if ri==0 else 'F3F3F1' if ri%2 else 'FFFFFF');pr.append(sh)
   margins=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    m=OxmlElement('w:'+side);m.set(qn('w:w'),'65');m.set(qn('w:type'),'dxa');margins.append(m)
   pr.append(margins)
  rp=cells[0]._tc.getparent().get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
  if ri==0:rp.append(OxmlElement('w:tblHeader'))
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(0)
 p.add_run().font.size=Pt(2)

def picture(path,page,description):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
 p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
 if page==11:
  # Native Word crop zooms the original panel; the screenshot pixels are unchanged.
  shape=p.add_run().add_picture(str(ROOT/path),width=Inches(5.8),height=Inches(2.93625))
  blipfill=shape._inline.graphic.graphicData.pic.blipFill
  rect=OxmlElement('a:srcRect');rect.set('l','25000');rect.set('r','25000');rect.set('t','20833');rect.set('b','34167');blipfill.insert(1,rect)
 else:
  width=6.7477 if page in [4,5] else 5.8
  shape=p.add_run().add_picture(str(ROOT/path),width=Inches(width))
 shape._inline.docPr.set('descr',description)

def cover(body):
 lines=body.splitlines();meta={}
 for line in lines:
  m=re.match(r'\*\*([^*]+):\*\*\s*(.+)',line)
  if m:meta[m[1]]=m[2]
 title=next(x[2:] for x in lines if x.startswith('# '));subtitle=next(x[3:] for x in lines if x.startswith('## '))
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(26);p.paragraph_format.space_after=Pt(75);p.paragraph_format.line_spacing=1
 logo=p.add_run().add_picture(str(ROOT/'assets/uts-logo.jpg'),width=Inches(5.416667))
 logo._inline.docPr.set('descr','University of Technology Sydney logo')
 def line(s,size,after=0,center=True):
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
  p.paragraph_format.space_after=Pt(after);p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True
  r=p.add_run(s);r.font.name='Liberation Serif';r.font.size=Pt(size);r.bold=True;r.font.color.rgb=RGBColor.from_string('202932')
  return p
 line(title,21,10);line(subtitle,16,22);line(meta['Subject'],14,140)
 line('Student Name: '+meta['Student Name'],12,5,False)
 p=line('Student ID: '+meta['Student ID'],12,18,False);p.paragraph_format.keep_with_next=False

chunks=re.split(r'<!-- PAGE (\d+) -->',source)
for idx in range(1,len(chunks),2):
 page=int(chunks[idx]);body=chunks[idx+1].strip()
 if page==1:cover(body);continue
 lines=body.splitlines();i=0;buf=[]
 def flush():
  if buf:para(' '.join(buf),page);buf.clear()
 while i<len(lines):
  line=lines[i].strip()
  if not line:flush();i+=1;continue
  if line.startswith('~~~'):
   flush();i+=1;code=[]
   while i<len(lines) and not lines[i].strip().startswith('~~~'):code.append(lines[i]);i+=1
   p=doc.add_paragraph();p.paragraph_format.line_spacing=1.03;p.paragraph_format.space_after=Pt(8)
   r=p.add_run('\n'.join(code));r.font.name='Liberation Mono';r.font.size=Pt(8.5)
   i+=1;continue
  if line.startswith('|'):
   flush();block=[]
   while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
   table(block,page);continue
  if line.startswith('# '):
   flush();p=doc.add_paragraph(line[2:],style='Heading 1');p.paragraph_format.page_break_before=True;i+=1;continue
  if line.startswith('## '):flush();doc.add_paragraph(line[3:],style='Heading 2');i+=1;continue
  m=re.fullmatch(r'!\[(.*)\]\((.*)\)',line)
  if m:flush();picture(m[2],page,m[1]);i+=1;continue
  if re.match(r'^\d+\. ',line):
   flush();p=para(line,page);p.paragraph_format.left_indent=Inches(.19);p.paragraph_format.first_line_indent=Inches(-.19);i+=1;continue
  if line.startswith('*Figure ') or line.startswith('*Table '):flush();para(line,page);i+=1;continue
  buf.append(line);i+=1
 flush()
# Remove inherited review markup and metadata residues from the delivered package.
for e in doc._element.xpath('.//w:commentRangeStart | .//w:commentRangeEnd | .//w:commentReference | .//w:vanish'):
 e.getparent().remove(e)
for e in doc.settings.element.xpath('./w:updateFields'):doc.settings.element.remove(e)
f=OxmlElement('w:updateFields');f.set(qn('w:val'),'true');doc.settings.element.append(f)

for style in doc.styles:
 if style.style_id.lower()=='title':
  for e in style.element.xpath('.//w:pBdr'):e.getparent().remove(e)
out=ROOT/'FormBridge-AT2-Report.docx';doc.save(out)
print(out)
print(json.dumps({'metrics':values,'main_pages':15,'appendix_pages':4},indent=2))
